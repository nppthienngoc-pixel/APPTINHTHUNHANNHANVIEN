from email.message import EmailMessage
from email.utils import formataddr
from io import BytesIO
import re
import smtplib
import socket
import ssl
import time
import zipfile

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

# =========================================================
# 1. CẤU HÌNH ỨNG DỤNG
# =========================================================

st.set_page_config(
    page_title="Quản lý lương và thuế TNCN", page_icon="💰", layout="wide"
)

st.title("💰 QUẢN LÝ LƯƠNG VÀ THUẾ TNCN - NGUYỄN CHÍ THANH")
st.caption(
    "Tính lương tối đa 500 nhân viên, xuất Excel và gửi phiếu lương riêng qua"
    " email."
)

st.warning(
    "Công thức thuế và mức giảm trừ được giữ theo code ban đầu. "
    "Hãy đối chiếu với quy định hiện hành trước khi dùng tính lương chính thức."
)

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DINH_DANG_SO = "#,##0.000;[Red]-#,##0.000;0.000"


# =========================================================
# 2. HÀM TÍNH THUẾ VÀ TIỆN ÍCH
# =========================================================


def tinh_thue_tncn(thu_nhap_tinh_thue):
    if thu_nhap_tinh_thue <= 0:
        return 0.0
    elif thu_nhap_tinh_thue <= 10:
        return thu_nhap_tinh_thue * 0.05
    elif thu_nhap_tinh_thue <= 30:
        return thu_nhap_tinh_thue * 0.10 - 0.5
    elif thu_nhap_tinh_thue <= 60:
        return thu_nhap_tinh_thue * 0.20 - 3.5
    elif thu_nhap_tinh_thue <= 100:
        return thu_nhap_tinh_thue * 0.30 - 9.5
    else:
        return thu_nhap_tinh_thue * 0.35 - 14.5


def dinh_dang_tien(x):
    return f"{x:,.3f}".replace(",", "X").replace(".", ",").replace("X", ".")


def email_hop_le(email):
    if not email:
        return False
    return bool(re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", str(email).strip()))


def lay_secret(ten, mac_dinh=""):
    """Đọc giá trị từ .streamlit/secrets.toml (nếu có)."""
    try:
        return st.secrets.get(ten, mac_dinh)
    except Exception:
        return mac_dinh


def tao_du_lieu_ban_dau(so_luong):
    return pd.DataFrame({
        "Mã nhân viên": [f"NV{i:03d}" for i in range(1, so_luong + 1)],
        "Họ và tên": ["" for _ in range(so_luong)],
        "Email": ["" for _ in range(so_luong)],
        "Lương cơ bản": [15.0 for _ in range(so_luong)],
        "Tiền thưởng": [0.0 for _ in range(so_luong)],
        "Phụ cấp tính BHXH": [0.0 for _ in range(so_luong)],
        "Phụ cấp không tính BHXH": [0.0 for _ in range(so_luong)],
        "Số người phụ thuộc": [0 for _ in range(so_luong)],
    })


# =========================================================
# 3. ĐỊNH DẠNG EXCEL
# =========================================================


def dinh_dang_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    ws.row_dimensions[1].height = 30

    for col in ws.columns:
        do_rong = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            do_rong + 3, 42
        )


def dinh_dang_cot_tien(ws, danh_sach_cot, cac_cot_tien):
    for ten_cot in cac_cot_tien:
        if ten_cot not in danh_sach_cot:
            continue
        vi_tri = danh_sach_cot.index(ten_cot) + 1
        for dong in range(2, ws.max_row + 1):
            ws.cell(row=dong, column=vi_tri).number_format = DINH_DANG_SO


# =========================================================
# 4. TẠO PHIẾU LƯƠNG CÁ NHÂN
# =========================================================


def tao_bao_cao_ca_nhan(nv):
    output = BytesIO()

    du_lieu = [
        ("Mã nhân viên", nv["Mã nhân viên"], ""),
        ("Họ và tên", nv["Họ và tên"], ""),
        ("Email", nv["Email"], ""),
        ("", "", ""),
        ("NỘI DUNG", "SỐ TIỀN", "ĐƠN VỊ"),
        ("Lương cơ bản", nv["Lương cơ bản"], "Triệu đồng"),
        ("Tiền thưởng", nv["Tiền thưởng"], "Triệu đồng"),
        ("Phụ cấp tính BHXH", nv["Phụ cấp tính BHXH"], "Triệu đồng"),
        ("Phụ cấp không tính BHXH", nv["Phụ cấp không tính BHXH"], "Triệu đồng"),
        ("TỔNG THU NHẬP", nv["Tổng thu nhập"], "Triệu đồng"),
        ("Tiền đóng BHXH", nv["Tiền đóng BHXH"], "Triệu đồng"),
        ("Giảm trừ bản thân", nv["Giảm trừ bản thân"], "Triệu đồng"),
        ("Số người phụ thuộc", nv["Số người phụ thuộc"], "Người"),
        (
            "Giảm trừ người phụ thuộc",
            nv["Giảm trừ người phụ thuộc"],
            "Triệu đồng",
        ),
        ("THU NHẬP TÍNH THUẾ", nv["Thu nhập tính thuế"], "Triệu đồng"),
        ("THUẾ TNCN", nv["Thuế TNCN"], "Triệu đồng"),
        ("TIỀN THỰC LÃNH", nv["Thực lãnh"], "Triệu đồng"),
    ]

    df = pd.DataFrame(du_lieu, columns=["Nội dung", "Giá trị", "Đơn vị"])

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Phieu luong", index=False)

        ws = writer.sheets["Phieu luong"]
        dinh_dang_sheet(ws)

        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 18

        for dong in [6, 7, 8, 9, 10, 11, 12, 14, 15, 16, 17]:
            ws.cell(row=dong, column=2).number_format = DINH_DANG_SO

        for dong in [10, 15, 16, 17]:
            for cell in ws[dong]:
                cell.font = Font(bold=True)

        dong_ghi_chu = ws.max_row + 2
        ws.cell(
            row=dong_ghi_chu,
            column=1,
            value=(
                "Đơn vị tiền: triệu đồng. Công thức tính theo mô hình được cài"
                " đặt trong ứng dụng."
            ),
        )
        ws.cell(
            row=dong_ghi_chu + 1,
            column=1,
            value=(
                "Cần kiểm tra quy định thuế và bảo hiểm hiện hành trước khi sử"
                " dụng chính thức."
            ),
        )

    output.seek(0)
    return output.getvalue()


# =========================================================
# 5. EXCEL TỔNG HỢP
# =========================================================


def tao_excel_tong_hop(bang):
    output = BytesIO()

    cac_cot = [
        "Mã nhân viên",
        "Họ và tên",
        "Email",
        "Lương cơ bản",
        "Tiền thưởng",
        "Phụ cấp tính BHXH",
        "Phụ cấp không tính BHXH",
        "Số người phụ thuộc",
        "Tổng thu nhập",
        "Tiền đóng BHXH",
        "Giảm trừ bản thân",
        "Giảm trừ người phụ thuộc",
        "Thu nhập tính thuế",
        "Thuế TNCN",
        "Thực lãnh",
    ]

    chi_tiet = bang[cac_cot].copy()

    tong_hop = pd.DataFrame({
        "Chỉ tiêu": [
            "Số người lao động",
            "Tổng thu nhập",
            "Tổng tiền đóng BHXH",
            "Tổng thu nhập tính thuế",
            "Tổng thuế TNCN",
            "Tổng tiền thực lãnh",
        ],
        "Giá trị": [
            len(bang),
            bang["Tổng thu nhập"].sum(),
            bang["Tiền đóng BHXH"].sum(),
            bang["Thu nhập tính thuế"].sum(),
            bang["Thuế TNCN"].sum(),
            bang["Thực lãnh"].sum(),
        ],
        "Đơn vị": [
            "Người",
            "Triệu đồng",
            "Triệu đồng",
            "Triệu đồng",
            "Triệu đồng",
            "Triệu đồng",
        ],
    })

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        chi_tiet.to_excel(writer, sheet_name="Bang luong chi tiet", index=False)
        tong_hop.to_excel(writer, sheet_name="Tong hop quy luong", index=False)

        dinh_dang_sheet(writer.sheets["Bang luong chi tiet"])
        dinh_dang_sheet(writer.sheets["Tong hop quy luong"])

        dinh_dang_cot_tien(
            writer.sheets["Bang luong chi tiet"],
            cac_cot,
            [
                "Lương cơ bản",
                "Tiền thưởng",
                "Phụ cấp tính BHXH",
                "Phụ cấp không tính BHXH",
                "Tổng thu nhập",
                "Tiền đóng BHXH",
                "Giảm trừ bản thân",
                "Giảm trừ người phụ thuộc",
                "Thu nhập tính thuế",
                "Thuế TNCN",
                "Thực lãnh",
            ],
        )

        ws = writer.sheets["Tong hop quy luong"]
        for dong in range(3, ws.max_row + 1):
            ws.cell(row=dong, column=2).number_format = DINH_DANG_SO

    output.seek(0)
    return output.getvalue()


# =========================================================
# 6. GỬI EMAIL (VIẾT LẠI)
# =========================================================


def chuan_hoa_mat_khau(mat_khau):
    """App Password Google hiển thị dạng 'abcd efgh ijkl mnop' -> bỏ mọi khoảng trắng."""
    return re.sub(r"\s+", "", str(mat_khau or ""))


def giai_thich_loi(e):
    """Chuyển lỗi kỹ thuật thành thông báo tiếng Việt dễ hiểu."""
    if isinstance(e, smtplib.SMTPAuthenticationError):
        return (
            "Đăng nhập thất bại (535). Nguyên nhân thường gặp: (1) chưa bật Xác"
            " minh 2 bước, (2) dùng mật khẩu Gmail thường thay vì App Password"
            " 16 ký tự, (3) App Password đã bị xóa hoặc gõ sai, (4) tài khoản"
            " Google Workspace bị quản trị viên tắt App Password / SMTP."
        )
    if isinstance(e, smtplib.SMTPRecipientsRefused):
        return "Máy chủ từ chối địa chỉ người nhận. Kiểm tra lại email nhân viên."
    if isinstance(e, smtplib.SMTPSenderRefused):
        return (
            "Máy chủ từ chối địa chỉ người gửi. Email gửi phải trùng với tài"
            " khoản đăng nhập."
        )
    if isinstance(e, smtplib.SMTPDataError):
        return (
            f"Máy chủ từ chối nội dung thư ({e.smtp_code}). Có thể đã vượt giới"
            " hạn gửi trong ngày của Gmail hoặc thư bị coi là spam."
        )
    if isinstance(e, smtplib.SMTPServerDisconnected):
        return "Máy chủ ngắt kết nối giữa chừng. Hãy thử lại."
    if isinstance(e, (socket.timeout, TimeoutError)):
        return (
            "Hết thời gian chờ. Mạng hoặc nơi chạy ứng dụng có thể đang chặn"
            " cổng SMTP. Thử đổi cổng 587 <-> 465."
        )
    if isinstance(e, socket.gaierror):
        return "Không tìm thấy máy chủ SMTP. Kiểm tra tên máy chủ và kết nối mạng."
    if isinstance(e, ssl.SSLError):
        return (
            "Lỗi SSL/TLS. Cổng 587 dùng STARTTLS, cổng 465 dùng SSL. Thử đổi"
            " cổng."
        )
    if isinstance(e, (ConnectionRefusedError, smtplib.SMTPConnectError, OSError)):
        return (
            f"Không kết nối được máy chủ SMTP ({e}). Mạng hoặc hosting có thể"
            " chặn cổng gửi mail."
        )
    return f"Lỗi: {e}"


def mo_ket_noi_smtp(smtp_server, smtp_port, email_gui, mat_khau):
    """Mở kết nối SMTP và đăng nhập. Cổng 465 dùng SSL, các cổng khác dùng STARTTLS."""
    context = ssl.create_default_context()
    port = int(smtp_port)

    if port == 465:
        server = smtplib.SMTP_SSL(smtp_server, port, context=context, timeout=30)
    else:
        server = smtplib.SMTP(smtp_server, port, timeout=30)

    try:
        if port != 465:
            server.ehlo()
            server.starttls(context=context)
            server.ehlo()
        server.login(email_gui, mat_khau)
    except Exception:
        try:
            server.close()
        except Exception:
            pass
        raise

    return server


def dong_ket_noi(server):
    try:
        server.quit()
    except Exception:
        try:
            server.close()
        except Exception:
            pass


def tao_tin_nhan(
    email_gui,
    ten_nguoi_gui,
    email_nhan,
    ho_ten,
    ma_nv,
    file_excel,
    tieu_de_mail,
    noi_dung_mail,
):
    ten_hien_thi = ho_ten if ho_ten else ma_nv

    def thay_the(text):
        return text.replace("{ho_ten}", ten_hien_thi).replace("{ma_nv}", ma_nv)

    msg = EmailMessage()
    # Xóa xuống dòng trong tiêu đề để tránh lỗi header
    msg["Subject"] = re.sub(r"[\r\n]+", " ", thay_the(tieu_de_mail)).strip()
    msg["From"] = formataddr((ten_nguoi_gui, email_gui))
    msg["To"] = email_nhan
    msg.set_content(thay_the(noi_dung_mail))

    msg.add_attachment(
        file_excel,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"Phieu_luong_{ma_nv}.xlsx",
    )
    return msg


def kiem_tra_ket_noi(smtp_server, smtp_port, email_gui, mat_khau):
    """Trả về (thành công, thông báo)."""
    try:
        server = mo_ket_noi_smtp(smtp_server, smtp_port, email_gui, mat_khau)
        dong_ket_noi(server)
        return True, "Kết nối và đăng nhập SMTP thành công."
    except Exception as e:
        return False, giai_thich_loi(e)


def gui_hang_loat(
    smtp_server,
    smtp_port,
    email_gui,
    ten_nguoi_gui,
    mat_khau,
    danh_sach,
    tieu_de_mail,
    noi_dung_mail,
    khoang_nghi,
    cap_nhat_tien_trinh,
):
    """Gửi tất cả bằng MỘT kết nối SMTP, tự kết nối lại nếu bị ngắt."""
    try:
        server = mo_ket_noi_smtp(smtp_server, smtp_port, email_gui, mat_khau)
    except Exception as e:
        return None, giai_thich_loi(e)

    ket_qua = []
    tong = len(danh_sach)

    try:
        for i, (_, nv) in enumerate(danh_sach.iterrows(), start=1):
            ma_nv = str(nv["Mã nhân viên"]).strip()
            email_nhan = str(nv["Email"]).strip()
            ho_ten = str(nv["Họ và tên"]).strip()

            try:
                msg = tao_tin_nhan(
                    email_gui,
                    ten_nguoi_gui,
                    email_nhan,
                    ho_ten,
                    ma_nv,
                    tao_bao_cao_ca_nhan(nv),
                    tieu_de_mail,
                    noi_dung_mail,
                )
                try:
                    server.send_message(msg)
                except smtplib.SMTPServerDisconnected:
                    server = mo_ket_noi_smtp(
                        smtp_server, smtp_port, email_gui, mat_khau
                    )
                    server.send_message(msg)

                ket_qua.append({
                    "Mã nhân viên": ma_nv,
                    "Họ và tên": ho_ten,
                    "Email": email_nhan,
                    "Kết quả": "Thành công",
                    "Chi tiết": "Email đã gửi thành công.",
                })
            except Exception as e:
                ket_qua.append({
                    "Mã nhân viên": ma_nv,
                    "Họ và tên": ho_ten,
                    "Email": email_nhan,
                    "Kết quả": "Thất bại",
                    "Chi tiết": giai_thich_loi(e),
                })

            cap_nhat_tien_trinh(i, tong)
            if i < tong:
                time.sleep(khoang_nghi)
    finally:
        dong_ket_noi(server)

    return ket_qua, None


# =========================================================
# 7. NHẬP DỮ LIỆU
# =========================================================

st.header("1. Nhập thông tin nhân viên")

so_luong = int(
    st.number_input(
        "Số người lao động cần tính",
        min_value=1,
        max_value=500,
        value=10,
        step=1,
    )
)

key_du_lieu = f"du_lieu_luong_{so_luong}"

if key_du_lieu not in st.session_state:
    st.session_state[key_du_lieu] = tao_du_lieu_ban_dau(so_luong)

du_lieu_nhap = st.data_editor(
    st.session_state[key_du_lieu],
    key=f"editor_luong_{so_luong}",
    num_rows="fixed",
    use_container_width=True,
    hide_index=True,
    height=min(600, 38 * (so_luong + 1)),
    column_config={
        "Mã nhân viên": st.column_config.TextColumn(
            "Mã nhân viên", required=True
        ),
        "Họ và tên": st.column_config.TextColumn("Họ và tên"),
        "Email": st.column_config.TextColumn("Email nhân viên"),
        "Lương cơ bản": st.column_config.NumberColumn(
            "Lương cơ bản (triệu)", min_value=0.0, step=0.5, format="%.3f"
        ),
        "Tiền thưởng": st.column_config.NumberColumn(
            "Tiền thưởng (triệu)", min_value=0.0, step=0.5, format="%.3f"
        ),
        "Phụ cấp tính BHXH": st.column_config.NumberColumn(
            "Phụ cấp tính BHXH (triệu)", min_value=0.0, step=0.5, format="%.3f"
        ),
        "Phụ cấp không tính BHXH": st.column_config.NumberColumn(
            "Phụ cấp không tính BHXH (triệu)",
            min_value=0.0,
            step=0.5,
            format="%.3f",
        ),
        "Số người phụ thuộc": st.column_config.NumberColumn(
            "Người phụ thuộc", min_value=0, max_value=100, step=1, format="%d"
        ),
    },
)


# =========================================================
# 8. TÍNH LƯƠNG
# =========================================================

if st.button("🧮 TÍNH LƯƠNG", type="primary", use_container_width=True):
    bang = du_lieu_nhap.copy()
    loi = []

    cot_tien = [
        "Lương cơ bản",
        "Tiền thưởng",
        "Phụ cấp tính BHXH",
        "Phụ cấp không tính BHXH",
    ]

    for cot in ["Mã nhân viên", "Họ và tên", "Email"]:
        bang[cot] = bang[cot].fillna("").astype(str).str.strip()

    if (bang["Mã nhân viên"] == "").any():
        loi.append("Mã nhân viên không được để trống.")

    if bang["Mã nhân viên"].duplicated().any():
        loi.append("Có mã nhân viên bị trùng.")

    for _, dong in bang.iterrows():
        if dong["Email"] and not email_hop_le(dong["Email"]):
            loi.append(f"Email không hợp lệ của nhân viên {dong['Mã nhân viên']}.")

    for cot in cot_tien:
        bang[cot] = pd.to_numeric(bang[cot], errors="coerce")

    bang["Số người phụ thuộc"] = pd.to_numeric(
        bang["Số người phụ thuộc"], errors="coerce"
    )

    if bang[cot_tien].isna().any().any():
        loi.append("Có khoản tiền trống hoặc không phải số.")

    if (bang[cot_tien] < 0).any().any():
        loi.append("Các khoản tiền không được âm.")

    if bang["Số người phụ thuộc"].isna().any():
        loi.append("Số người phụ thuộc không được để trống.")
    elif (
        (bang["Số người phụ thuộc"] < 0) | (bang["Số người phụ thuộc"] % 1 != 0)
    ).any():
        loi.append("Số người phụ thuộc phải là số nguyên không âm.")

    if loi:
        for thong_bao in loi:
            st.error(thong_bao)
    else:
        bang["Số người phụ thuộc"] = bang["Số người phụ thuộc"].astype(int)

        bang["Tổng thu nhập"] = (
            bang["Lương cơ bản"]
            + bang["Tiền thưởng"]
            + bang["Phụ cấp tính BHXH"]
            + bang["Phụ cấp không tính BHXH"]
        )
        bang["Tiền đóng BHXH"] = (
            bang["Lương cơ bản"] + bang["Phụ cấp tính BHXH"]
        ) * 0.105
        bang["Giảm trừ bản thân"] = 15.5
        bang["Giảm trừ người phụ thuộc"] = bang["Số người phụ thuộc"] * 6.2
        bang["Thu nhập tính thuế"] = (
            bang["Tổng thu nhập"]
            - bang["Tiền đóng BHXH"]
            - bang["Giảm trừ bản thân"]
            - bang["Giảm trừ người phụ thuộc"]
        )
        bang["Thuế TNCN"] = bang["Thu nhập tính thuế"].apply(tinh_thue_tncn)
        bang["Thực lãnh"] = (
            bang["Tổng thu nhập"] - bang["Tiền đóng BHXH"] - bang["Thuế TNCN"]
        )

        st.session_state["ket_qua_bang_luong"] = bang.copy()
        st.session_state.pop("ket_qua_gui_email", None)
        st.success(f"Đã tính lương cho {len(bang)} nhân viên.")


# =========================================================
# 9. HIỂN THỊ KẾT QUẢ VÀ XUẤT FILE
# =========================================================

if "ket_qua_bang_luong" in st.session_state:
    bang_kq = st.session_state["ket_qua_bang_luong"]

    st.divider()
    st.header("2. Bảng lương kết quả")

    cot_hien_thi = [
        "Mã nhân viên",
        "Họ và tên",
        "Email",
        "Tổng thu nhập",
        "Tiền đóng BHXH",
        "Giảm trừ bản thân",
        "Giảm trừ người phụ thuộc",
        "Thu nhập tính thuế",
        "Thuế TNCN",
        "Thực lãnh",
    ]

    st.dataframe(bang_kq[cot_hien_thi], use_container_width=True, hide_index=True)

    st.subheader("Tổng hợp quỹ lương")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Số nhân viên", f"{len(bang_kq):,}")
    c2.metric(
        "Tổng thu nhập", f"{dinh_dang_tien(bang_kq['Tổng thu nhập'].sum())} triệu"
    )
    c3.metric(
        "Tổng thuế TNCN", f"{dinh_dang_tien(bang_kq['Thuế TNCN'].sum())} triệu"
    )
    c4.metric(
        "Tổng thực lãnh", f"{dinh_dang_tien(bang_kq['Thực lãnh'].sum())} triệu"
    )

    st.subheader("3. Tải báo cáo Excel")

    st.download_button(
        "📥 Tải Excel tổng hợp toàn bộ nhân viên",
        data=tao_excel_tong_hop(bang_kq),
        file_name="Bang_luong_tong_hop.xlsx",
        mime=MIME_XLSX,
        use_container_width=True,
    )

    zip_output = BytesIO()

    with zipfile.ZipFile(
        zip_output, "w", compression=zipfile.ZIP_DEFLATED
    ) as zip_file:
        ten_da_dung = set()

        for _, nv in bang_kq.iterrows():
            ma_nv = (
                re.sub(r'[\\/*?:"<>|]', "_", str(nv["Mã nhân viên"]).strip()).strip(
                    " ._"
                )
                or "Nhan_vien"
            )

            ten_goc = ma_nv
            dem = 2
            while f"Phieu_luong_{ma_nv}.xlsx" in ten_da_dung:
                ma_nv = f"{ten_goc}_{dem}"
                dem += 1

            ten_file = f"Phieu_luong_{ma_nv}.xlsx"
            ten_da_dung.add(ten_file)
            zip_file.writestr(ten_file, tao_bao_cao_ca_nhan(nv))

    zip_output.seek(0)

    st.download_button(
        "📦 Tải ZIP phiếu lương riêng của từng nhân viên",
        data=zip_output.getvalue(),
        file_name="Phieu_luong_tung_nhan_vien.zip",
        mime="application/zip",
        use_container_width=True,
    )

    st.subheader("4. Xem và tải phiếu lương cá nhân")

    danh_sach_ma = bang_kq["Mã nhân viên"].astype(str).tolist()
    ma_chon = st.selectbox("Chọn mã nhân viên", options=danh_sach_ma)
    nv_chon = bang_kq[bang_kq["Mã nhân viên"].astype(str) == ma_chon].iloc[0]

    st.write(f"**Họ tên:** {nv_chon['Họ và tên'] or 'Chưa nhập'}")
    st.write(f"**Email:** {nv_chon['Email'] or 'Chưa nhập'}")

    st.download_button(
        f"📄 Tải phiếu lương của {ma_chon}",
        data=tao_bao_cao_ca_nhan(nv_chon),
        file_name=f"Phieu_luong_{ma_chon}.xlsx",
        mime=MIME_XLSX,
        use_container_width=True,
    )


# =========================================================
# 10. GỬI PHIẾU LƯƠNG QUA EMAIL
# =========================================================

if "ket_qua_bang_luong" in st.session_state:
    bang_gui = st.session_state["ket_qua_bang_luong"]

    st.divider()
    st.header("5. Gửi phiếu lương qua email")

    with st.expander("📖 Hướng dẫn tạo App Password để gửi được Gmail", expanded=False):
        st.markdown(
            """
**Gmail không cho đăng nhập SMTP bằng mật khẩu thường.** Bạn cần:

1. Đăng nhập tài khoản Gmail gửi thư, vào **myaccount.google.com → Bảo mật**.
2. Bật **Xác minh 2 bước**.
3. Vào **Mật khẩu ứng dụng** (App passwords), tạo mật khẩu mới, đặt tên bất kỳ.
4. Sao chép chuỗi 16 ký tự Google hiển thị (có thể có khoảng trắng, app sẽ tự bỏ).

**Nếu là Gmail công ty (Google Workspace)** mà không thấy mục "Mật khẩu ứng dụng":
quản trị viên Workspace cần cho phép (Admin Console → Bảo mật → Xác thực → Xác minh 2 bước,
và bật quyền truy cập SMTP/ứng dụng kém bảo mật hơn cho người dùng).

**Mẹo bảo mật:** thay vì gõ mật khẩu mỗi lần, tạo file `.streamlit/secrets.toml`:

```toml
EMAIL_CONG_TY = "email_cua_ban@gmail.com"
APP_PASSWORD = "abcdefghijklmnop"
```
"""
        )

    col_a, col_b = st.columns(2)

    with col_a:
        email_cong_ty = st.text_input(
            "Email gửi thư (tài khoản Gmail)",
            value=lay_secret("EMAIL_CONG_TY", "nppthienngoc@gmail.com"),
        ).strip()

    with col_b:
        ten_nguoi_gui = st.text_input("Tên hiển thị người gửi", value="Phòng Nhân sự")

    col_smtp1, col_smtp2 = st.columns([3, 1])

    with col_smtp1:
        smtp_server = st.text_input("Máy chủ SMTP", value="smtp.gmail.com").strip()

    with col_smtp2:
        smtp_port = st.selectbox(
            "Cổng SMTP",
            options=[587, 465],
            help="587 = STARTTLS (khuyên dùng), 465 = SSL. Nếu lỗi kết nối, thử cổng còn lại.",
        )

    mat_khau_secret = lay_secret("APP_PASSWORD", "")
    if mat_khau_secret:
        st.info("Đang dùng App Password từ file secrets.toml.")
        mat_khau = chuan_hoa_mat_khau(mat_khau_secret)
    else:
        mat_khau = chuan_hoa_mat_khau(
            st.text_input(
                "Mật khẩu ứng dụng (App Password) 16 ký tự",
                type="password",
            )
        )

    # ---- Kiểm tra kết nối ----
    if st.button("🔌 Kiểm tra kết nối SMTP"):
        if not email_hop_le(email_cong_ty):
            st.error("Email gửi thư không hợp lệ.")
        elif not mat_khau:
            st.error("Vui lòng nhập App Password.")
        else:
            with st.spinner("Đang kết nối tới máy chủ SMTP..."):
                ok, thong_diep = kiem_tra_ket_noi(
                    smtp_server, smtp_port, email_cong_ty, mat_khau
                )
            if ok:
                st.success(thong_diep)
            else:
                st.error(thong_diep)

    st.subheader("Nội dung email gửi nhân viên")

    tieu_de_mail = st.text_input(
        "Tiêu đề email",
        value="Phiếu lương cá nhân - {ma_nv}",
        help="Dùng {ho_ten} hoặc {ma_nv} để tự động thay thế.",
    )

    noi_dung_mac_dinh = """Kính gửi {ho_ten},

Phòng nhân sự gửi bạn phiếu lương cá nhân đính kèm.

Vui lòng kiểm tra các khoản thu nhập, bảo hiểm, thuế TNCN và số tiền thực lãnh. Nếu có sai sót, vui lòng phản hồi lại email này.

Đề nghị bảo mật thông tin trong file đính kèm.

Trân trọng,
Phòng Nhân sự"""

    noi_dung_mail = st.text_area(
        "Nội dung email",
        value=noi_dung_mac_dinh,
        height=200,
        help="Dùng {ho_ten} hoặc {ma_nv} để tự động thay thế.",
    )

    # Danh sách nhân viên có email hợp lệ
    danh_sach_nhan = bang_gui[
        bang_gui["Email"].fillna("").astype(str).str.strip().apply(email_hop_le)
    ].copy()

    so_thieu = len(bang_gui) - len(danh_sach_nhan)
    st.write(
        f"👥 **{len(danh_sach_nhan)}** nhân viên có email hợp lệ"
        + (f" ({so_thieu} người chưa có email, sẽ bị bỏ qua)." if so_thieu else ".")
    )

    pham_vi_gui = st.radio(
        "Chọn người nhận",
        [
            "Gửi cho một nhân viên (nên thử trước)",
            "Gửi cho tất cả nhân viên có email hợp lệ",
        ],
    )

    ma_nv_chon = None
    if pham_vi_gui.startswith("Gửi cho một"):
        if not danh_sach_nhan.empty:
            ds_ma = danh_sach_nhan["Mã nhân viên"].astype(str).tolist()
            nhan_hien_thi = {
                str(r["Mã nhân viên"]): (
                    f"{r['Mã nhân viên']} - {r['Họ và tên']} ({r['Email']})"
                )
                for _, r in danh_sach_nhan.iterrows()
            }
            ma_nv_chon = st.selectbox(
                "Chọn nhân viên nhận phiếu lương",
                options=ds_ma,
                format_func=lambda m: nhan_hien_thi.get(m, m),
            )
        else:
            st.warning("Chưa có email nhân viên hợp lệ trong bảng.")

    khoang_nghi = st.slider(
        "Khoảng nghỉ giữa các email (giây)",
        min_value=0.0,
        max_value=5.0,
        value=1.0,
        step=0.5,
        help="Gửi quá nhanh dễ bị Gmail giới hạn. Gmail thường cho khoảng 500 thư/ngày (tài khoản cá nhân).",
    )

    xac_nhan_gui = st.checkbox(
        "Tôi đã kiểm tra danh sách người nhận và xác nhận "
        "gửi đúng phiếu lương cho từng nhân viên."
    )

    bam_gui = st.button("📨 GỬI PHIẾU LƯƠNG", type="primary", use_container_width=True)

    if bam_gui:
        if not email_hop_le(email_cong_ty):
            st.error("Email gửi thư không hợp lệ.")
        elif not mat_khau:
            st.error("Vui lòng nhập Mật khẩu ứng dụng (App Password).")
        elif not xac_nhan_gui:
            st.error("Bạn phải tích chọn ô xác nhận danh sách người nhận.")
        elif danh_sach_nhan.empty:
            st.error("Không có nhân viên nào có địa chỉ email hợp lệ.")
        else:
            if pham_vi_gui.startswith("Gửi cho một"):
                danh_sach_can_gui = danh_sach_nhan[
                    danh_sach_nhan["Mã nhân viên"].astype(str) == ma_nv_chon
                ]
            else:
                danh_sach_can_gui = danh_sach_nhan

            tien_trinh = st.progress(0)
            thong_bao = st.empty()

            def cap_nhat_tien_trinh(i, tong):
                tien_trinh.progress(i / tong)
                thong_bao.write(f"Đã xử lý {i}/{tong} email.")

            with st.spinner("Đang kết nối và gửi email..."):
                ket_qua_gui, loi_ket_noi = gui_hang_loat(
                    smtp_server=smtp_server,
                    smtp_port=smtp_port,
                    email_gui=email_cong_ty,
                    ten_nguoi_gui=ten_nguoi_gui,
                    mat_khau=mat_khau,
                    danh_sach=danh_sach_can_gui,
                    tieu_de_mail=tieu_de_mail,
                    noi_dung_mail=noi_dung_mail,
                    khoang_nghi=khoang_nghi,
                    cap_nhat_tien_trinh=cap_nhat_tien_trinh,
                )

            if loi_ket_noi:
                st.error(f"Không thể kết nối/đăng nhập SMTP. {loi_ket_noi}")
            else:
                df_gui = pd.DataFrame(ket_qua_gui)
                st.session_state["ket_qua_gui_email"] = df_gui

                thanh_cong = int((df_gui["Kết quả"] == "Thành công").sum())
                that_bai = len(df_gui) - thanh_cong

                if thanh_cong > 0:
                    st.success(f"Hoàn tất: {thanh_cong} thành công, {that_bai} thất bại.")
                else:
                    st.error("Gửi email thất bại hoàn toàn. Xem chi tiết lỗi bên dưới.")

    if "ket_qua_gui_email" in st.session_state:
        st.subheader("Bảng chi tiết kết quả gửi email")

        st.dataframe(
            st.session_state["ket_qua_gui_email"],
            use_container_width=True,
            hide_index=True,
        )

        st.download_button(
            "📥 Tải báo cáo trạng thái gửi email",
            data=st.session_state["ket_qua_gui_email"]
            .to_csv(index=False)
            .encode("utf-8-sig"),
            file_name="Ket_qua_gui_email.csv",
            mime="text/csv",
        )


st.caption(
    "Bảo mật: Phiếu lương chứa dữ liệu cá nhân. "
    "Chỉ gửi đúng phiếu lương cho người nhận tương ứng."
)
